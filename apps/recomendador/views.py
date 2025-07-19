from django.shortcuts import render, redirect
from .forms import ComercioForm
from .models import Comercio
from .recommender import RecomendadorEmpresas
from django.contrib import messages
import os
import openpyxl

# Ruta absoluta segura para producción y Railway
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, 'apps', 'recomendador', 'data', 'base_actualizada.xlsx')

# Recomendador global
recomendador = None

# Función para cargar o recargar el modelo
def cargar_recomendador():
    global recomendador
    if os.path.exists(EXCEL_PATH):
        recomendador = RecomendadorEmpresas(EXCEL_PATH)

# Carga inicial
cargar_recomendador()


def index(request):
    recomendaciones = None

    if request.method == 'POST':
        consulta = request.POST.get('consulta', '')
        if consulta and recomendador:
            recomendaciones = recomendador.recomendar(consulta).to_dict(orient='records')
            if not recomendaciones:
                messages.warning(request, 'No se encontraron resultados para tu búsqueda.')
        else:
            messages.error(request, 'No se pudo generar recomendaciones. Verifica la base o el texto ingresado.')

    return render(request, 'recomendador/index.html', {'recomendaciones': recomendaciones})


def registrar_comercio(request):
    if request.method == 'POST':
        form = ComercioForm(request.POST, request.FILES)
        if form.is_valid():
            comercio = form.save()

            # Verificar si existe el archivo y crear estructura si no
            if not os.path.exists(EXCEL_PATH):
                os.makedirs(os.path.dirname(EXCEL_PATH), exist_ok=True)
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = 'BBDD'
                ws.append([
                    'NOMBRE', 'SECTOR', 'SUBSECTOR', 'ARTICULOS',
                    'DIRECCIÓN', 'CELULAR', 'TELÉFONO',
                    'FACEBOOK', 'INSTAGRAM'
                ])
            else:
                wb = openpyxl.load_workbook(EXCEL_PATH)
                ws = wb['BBDD'] if 'BBDD' in wb.sheetnames else wb.active

            ws.append([
                comercio.nombre,
                comercio.sector,
                comercio.subsector,
                comercio.articulos,
                comercio.direccion,
                comercio.celular,
                comercio.telefono,
                comercio.link_facebook,
                comercio.link_instagram,
            ])

            wb.save(EXCEL_PATH)

            # Recargar el recomendador después de registrar
            cargar_recomendador()

            messages.success(request, '¡Comercio registrado exitosamente!')
            return redirect('registro')
        else:
            messages.error(request, 'Error en el formulario. Por favor revisa los campos.')
    else:
        form = ComercioForm()

    return render(request, 'recomendador/registro.html', {'form': form})
